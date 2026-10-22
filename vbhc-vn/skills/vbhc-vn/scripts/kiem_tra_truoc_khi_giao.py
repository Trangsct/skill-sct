#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""kiem_tra_truoc_khi_giao.py — MỘT LỆNH kiểm tra trước khi giao file (Bạn chốt 17/10/2026, bản giao việc mục 3.6).

Gộp, chạy một lần:
  1. qa_all.py        — thể thức, nội dung, render PDF, quy tắc R01–R23 (R23: số hiệu văn bản quy phạm có nguồn — GATE 0)
  2. cite_check.py    — số, ngày văn bản viện dẫn khớp sổ đăng ký (LỆCH là lỗi)
  3. ra_soat_doc_lap.py — số đo cho GATE cuối, dùng lại PDF qa_all.py vừa render (không render lần hai)

    python3 scripts/kiem_tra_truoc_khi_giao.py file.docx
    python3 scripts/kiem_tra_truoc_khi_giao.py file.docx --tu-sua      # tự sửa lỗi hình thức rồi kiểm lại (tối đa 2 vòng)
    python3 scripts/kiem_tra_truoc_khi_giao.py file.docx --final       # bản Bạn yêu cầu hoàn thiện để xuất bản
    python3 scripts/kiem_tra_truoc_khi_giao.py file.docx --anh         # xuất thêm ảnh ghép soi trang (chậm hơn)

--tu-sua (lỗi máy sửa được, KHÔNG hỏi Bạn): chuan_hoa_xuat_ban.py (phông, Nơi nhận cỡ 11, một dòng trước khối ký,
văn bản thủ tục hành chính bỏ Ban Giám đốc Sở); đúng một dòng trống sau khối "Kính gửi"; bảng nội dung nền trắng, căn
giữa theo chiều dọc; dòng Lưu bỏ tên chuyên viên. Sửa TẠI CHỖ file đưa vào. Còn lỗi sau 2 vòng thì in ra để bôi tím và báo Bạn, không kéo dài.

Mục tiêu thời gian dưới 30 giây trên docx thường; dòng cuối in thời gian thực đo.
Thoát mã 0 = không FAIL; 1 = có FAIL (qa_all.py) hoặc LỆCH số, ngày (cite_check.py); 2 = lỗi đọc file.
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
import time
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

import ra_soat_doc_lap  # noqa: E402
from cite_check import quet as cite_quet  # noqa: E402
from qa_rules import Ctx, _la_tthc, _o_noi_nhan, _tt  # noqa: E402

from qa_all import QA_DIR as QA_PDF_DIR  # noqa: E402
HINH_THUC = ("1. Phông", "2. Nơi nhận", "3. Sau Kính gửi", "3. Trước khối ký", "7. Bảng")


def _trong(p_el) -> bool:
    return (not _tt(p_el).strip() and p_el.find(".//" + qn("w:drawing")) is None
            and p_el.find(".//" + qn("w:pict")) is None and p_el.find(".//" + qn("w:sectPr")) is None
            and p_el.find(".//" + qn("w:br")) is None)


def sua_kinh_gui(doc) -> int:
    """Đúng một đoạn trống sau khối Kính gửi (đoạn thân, không áp cho bảng Kính gửi)."""
    body = doc.element.body
    kids = list(body)
    for i, el in enumerate(kids):
        if el.tag != qn("w:p") or not _tt(el).strip().startswith("Kính gửi"):
            continue
        j = i + 1
        while j < len(kids) and kids[j].tag == qn("w:p") and _tt(kids[j]).strip().startswith(("-", "+")):
            j += 1
        cuoi = kids[j - 1]
        trong = []
        while j < len(kids) and kids[j].tag == qn("w:p") and _trong(kids[j]):
            trong.append(kids[j])
            j += 1
        if len(trong) == 1:
            return 0
        if not trong:
            p = OxmlElement("w:p")
            cuoi.addnext(p)
            return 1
        for el2 in trong[1:]:
            body.remove(el2)
        return len(trong) - 1
    return 0


def sua_bang(doc) -> int:
    """Bảng nội dung (trừ bảng đầu trang, bảng Nơi nhận - ký, bảng Kính gửi): nền trắng, căn giữa chiều dọc."""
    t_ky, _ = _o_noi_nhan(doc)
    n = 0
    for ti, tb in enumerate(doc.tables):
        if ti == 0 or (t_ky is not None and tb._tbl is t_ky._tbl) or "Kính gửi" in tb._tbl.xpath("string(.)"):
            # so sánh phần tử XML: doc.tables tạo đối tượng Table mới mỗi lần gọi nên "tb is t_ky" luôn sai (lỗi 22/10/2026)
            continue
        for tc in tb._tbl.iter(qn("w:tc")):
            pr = tc.get_or_add_tcPr()
            shd = pr.find(qn("w:shd"))
            if shd is not None and (shd.get(qn("w:fill")) or "").upper() not in ("", "FFFFFF", "AUTO"):
                pr.remove(shd)
                n += 1
            va = pr.find(qn("w:vAlign"))
            if va is None or va.get(qn("w:val")) != "center":
                if va is None:
                    va = OxmlElement("w:vAlign")
                    pr.append(va)
                va.set(qn("w:val"), "center")
                n += 1
    return n


def sua_dong_luu(doc) -> int:
    """Dòng Lưu không ghi tên chuyên viên (Bạn chốt 01/10/2026, R07): 'Lưu: VT, CN(Linh).' → 'Lưu: VT, CN.'."""
    _, o = _o_noi_nhan(doc)
    if o is None:
        return 0
    n = 0
    for p in o.paragraphs:
        if re.match(r"^\s*-?\s*Lưu\s*:", p.text):
            cu = p.text
            moi = re.sub(r"\s*\((?!b/c\b)[^()]{1,30}\)", "", cu)
            if moi != cu:  # tên chuyên viên thường tách nhiều run — ghép vào run có chữ đầu tiên
                co_chu = [r for r in p.runs if r.text.strip()]
                for r in co_chu[1:]:
                    r.text = ""
                co_chu[0].text = moi.lstrip() if not cu[:1].isspace() else moi
                n += 1
    return n


def tu_sua(path: Path) -> list[str]:
    from chuan_hoa_xuat_ban import chuan_hoa_file
    doc0 = Document(str(path))
    ctx = Ctx(path)
    ctx.build(doc0)
    tthc = _la_tthc(ctx)
    chuan_hoa_file(path, None, tthc=tthc)
    doc = Document(str(path))
    a = sua_kinh_gui(doc)
    b = sua_bang(doc)
    c = sua_dong_luu(doc)
    doc.save(str(path))
    if a or b or c:  # python-docx lưu lại có thể sinh phông mặc định — chuẩn hóa thêm một lần
        chuan_hoa_file(path, None, tthc=tthc)
    return [f"chuan_hoa_xuat_ban.py{' --tthc' if tthc else ''}", f"dòng trống sau Kính gửi: {a} chỗ",
            f"ô bảng nền trắng, căn giữa dọc: {b} thuộc tính", f"bỏ tên chuyên viên ở dòng Lưu: {c} chỗ"]


def chay_qa_all(path: Path, final: bool, anh: bool) -> tuple[int, list[str]]:
    lenh = [sys.executable, str(SCRIPT_DIR / "qa_all.py"), str(path)]
    if final:
        lenh.append("--final")
    if not anh:
        lenh.append("--no-image")
    r = subprocess.run(lenh, capture_output=True, text=True, timeout=300)
    dong, dem = [], {}
    for ln in r.stdout.splitlines():
        if not ln.startswith(("[FAIL", "[WARN", "KẾT QUẢ", "ẢNH GHÉP")):
            continue
        tag = ln.split("]", 1)[0]
        dem[tag] = dem.get(tag, 0) + 1
        if dem[tag] <= 3 or not ln.startswith("["):  # cùng mã quá 3 dòng thì gộp (vd R03 trên bản nhiều chữ tím)
            dong.append(ln)
    dong += [f"{tag}] … thêm {n - 3} dòng cùng mã" for tag, n in dem.items() if n > 3]
    return r.returncode, dong


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("--tu-sua", action="store_true")
    ap.add_argument("--final", action="store_true")
    ap.add_argument("--anh", action="store_true")
    a = ap.parse_args()
    path = Path(a.docx)
    if not path.exists():
        print(f"LỖI: không tìm thấy {path}", file=sys.stderr)
        return 2
    t0 = time.time()
    vong = 2 if a.tu_sua else 1
    for v in range(1, vong + 1):
        if a.tu_sua:
            for x in tu_sua(path):
                print(f"[TỰ SỬA vòng {v}] {x}")
        rc_qa, dong_qa = chay_qa_all(path, a.final, a.anh)
        pdf = QA_PDF_DIR / (path.stem + ".pdf")
        do = ra_soat_doc_lap.do(path, pdf if pdf.exists() else None, render=False)
        con_hinh_thuc = [x for x in do if x["canh_bao"] and x["muc"].startswith(HINH_THUC)]
        if not (a.tu_sua and con_hinh_thuc):
            break
    cite = cite_quet(path)
    lech = [x for x in cite if x["trang_thai"] == "LỆCH"]

    print("=" * 66)
    print(f"KIỂM TRA TRƯỚC KHI GIAO — {path.name}")
    print("=" * 66)
    print("── 1. qa_all.py ──")
    for ln in dong_qa:
        print("  " + ln)
    print("── 2. cite_check.py ──")
    print(f"  {len(cite)} số hiệu: khớp {sum(1 for x in cite if x['trang_thai'] == 'KHỚP')}, lệch {len(lech)}, "
          f"chưa có trong sổ đăng ký {sum(1 for x in cite if x['trang_thai'] == 'CHƯA CÓ')}")
    for x in lech:
        print(f"  [LỆCH] {x['so_hieu']} — {x['ghi_chu']}")
    print("── 3. ra_soat_doc_lap.py (số đo cho GATE cuối; đối chiếu checklist_ra_soat.txt) ──")
    for x in do:
        if x["canh_bao"]:
            print(f"  ⚠ [{x['muc']}] {x['noi_dung']}")
    n_ht = sum(1 for x in do if x["canh_bao"] and x["muc"].startswith(HINH_THUC))
    n_nd = sum(1 for x in do if x["canh_bao"]) - n_ht
    print(f"  {n_ht} chỗ hình thức, {n_nd} chỗ nội dung cần xem")
    loi = rc_qa != 0 or bool(lech)
    print(f"\nKẾT QUẢ: {'FAIL' if loi else 'PASS'}  |  thời gian: {time.time() - t0:.1f}s (mục tiêu < 30s)")
    return 1 if loi else 0


if __name__ == "__main__":
    sys.exit(main())
