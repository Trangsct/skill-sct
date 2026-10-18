#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
qa_pdf_check.py — QA tự động thể thức VBHC trước khi giao file.

Sinh ra từ vụ 4 lỗi thật ngày 05/7/2026 (công văn HHNH loại 5, 8):
  1. WIDOW    — 1 chữ đơn độc rơi xuống dòng cuối đoạn (quét pdftotext -layout).
  2. SIGSPLIT — khối chữ ký gãy giữa 2 trang: tên người ký khác trang với
                chức danh "KT. GIÁM ĐỐC"/"TM. .../Q. GIÁM ĐỐC/GIÁM ĐỐC".
  3. LINES    — thiếu đường Line header (văn bản SCT chuẩn phải có >= 2
                shape w:pict/w:drawing: dưới tên cơ quan và dưới Tiêu ngữ).
  4. SZ13     — dòng "Số: .../..." và "…, ngày … tháng … năm …" chưa đặt
                13pt tường minh (w:sz=26) hoặc dòng ngày thiếu in nghiêng.

Cách dùng:
  python3 qa_pdf_check.py <file.docx> [--min-lines 2] [--no-render] [--pdf <pdf-đã-render>]

Thoát mã 0 = PASS toàn bộ; 1 = có FAIL. Cảnh báo (WARN) không chặn.
Yêu cầu: soffice (LibreOffice), pdftotext, python-docx.
"""

import re
import unicodedata
import subprocess
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    print("Cần python-docx: pip install python-docx --break-system-packages")
    sys.exit(2)

SIGN_TITLES = [
    "KT. GIÁM ĐỐC", "PHÓ GIÁM ĐỐC", "GIÁM ĐỐC", "TM. ỦY BAN NHÂN DÂN",
    "CHỦ TỊCH", "KT. CHỦ TỊCH", "PHÓ CHỦ TỊCH", "TL. CHỦ TỊCH",
    "CHÁNH VĂN PHÒNG", "KT. CHÁNH VĂN PHÒNG", "PHÓ CHÁNH VĂN PHÒNG",
    "TRƯỞNG PHÒNG", "TM. ĐOÀN KIỂM TRA", "TRƯỞNG ĐOÀN", "Q. GIÁM ĐỐC",
]

# Các dòng 1 từ hợp lệ, không tính widow (số trang, ký hiệu...)
WIDOW_WHITELIST = re.compile(r"^(\d+|[IVXLC]+\.?|PHỤ|LỤC|-|–|_{3,}|\.{3,})$")


def render_pdf(docx_path: Path, outdir: Path) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf",
         "--outdir", str(outdir), str(docx_path)],
        check=True, capture_output=True, timeout=120,
    )
    pdf = outdir / (docx_path.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("Render PDF thất bại")
    return pdf


def pdf_pages_text(pdf_path: Path):
    out = subprocess.run(
        ["pdftotext", "-layout", str(pdf_path), "-"],
        check=True, capture_output=True, timeout=60,
    ).stdout.decode("utf-8", errors="replace")
    return out.split("\f")


def check_widow(pages) -> list:
    """Dòng chỉ có 1 từ đứng ngay trước dòng trống (cuối đoạn)."""
    issues = []
    for pno, page in enumerate(pages, 1):
        lines = page.split("\n")
        for i, ln in enumerate(lines):
            s = ln.strip()
            if not s or len(s.split()) != 1 or WIDOW_WHITELIST.match(s):
                continue
            nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
            prv = lines[i - 1].strip() if i > 0 else ""
            # widow = 1 từ, cuối đoạn (dòng sau trống), dòng trước có nội dung dài
            if nxt == "" and len(prv.split()) >= 4:
                issues.append(f"trang {pno}: '{s}' rơi 1 mình cuối đoạn (sau: '{prv[-50:]}')")
    return issues


def _gon(s: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", s or "")).strip()


def _o_ky_docx(docx_path: Path):
    """Các dòng chữ của KHỐI KÝ cuối cùng trong file .docx (chức danh đầu khối, tên người ký cuối khối).

    Ưu tiên ô bảng chứa chức danh ký (bảng Nơi nhận - chữ ký, ô ký của giấy chứng nhận, biên bản); không có bảng thì
    lấy đoạn thân cuối cùng chứa chức danh ký và tối đa 8 đoạn sau nó. Không thấy → None."""
    d = Document(str(docx_path))
    # dòng chức danh của khối ký: ngắn, viết hoa toàn bộ ("KT. GIÁM ĐỐC", "PHÓ CHỦ TỊCH") — không phải câu thân văn bản
    co_chuc_danh = lambda x: len(x) <= 40 and x == x.upper() and any(tt in x for tt in SIGN_TITLES)  # noqa: E731
    o_cuoi = None
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                dong = [_gon(p.text) for p in c.paragraphs if p.text.strip()]
                if len(dong) >= 2 and co_chuc_danh(dong[0]) and not co_chuc_danh(dong[-1]):
                    o_cuoi = dong
    if o_cuoi:
        return o_cuoi
    ps = [_gon(p.text) for p in d.paragraphs]
    vt = max((i for i, x in enumerate(ps) if x and co_chuc_danh(x)), default=None)
    if vt is None:
        return None
    dong = [x for x in ps[vt:vt + 9] if x]
    while len(dong) > 1 and co_chuc_danh(dong[-1]):
        dong.pop()
    return dong if len(dong) >= 2 else None


def check_sig_split(pages, docx_path: Path | None = None) -> list:
    """Chức danh ký và tên người ký phải cùng trang.

    Có file .docx (cách chuẩn từ 17/10/2026): lấy đúng các dòng của ô ký trong bảng Nơi nhận - chữ ký (chức danh
    đầu ô, tên người ký cuối ô), tìm trang chứa tên người ký (lần xuất hiện cuối), rồi trang chứa dòng chức danh
    đầu ô gần nhất phía trước. Khác trang → gãy khối ký. Trang SAU khối ký (phụ lục, biểu kèm theo) không tính.
    Vụ 17/10/2026: cách cũ báo nhầm mọi văn bản có phụ lục sau khối ký, và bỏ sót khi "KT. GIÁM ĐỐC" ở trang 1
    còn "PHÓ GIÁM ĐỐC" + tên sang trang 2.

    Không có .docx hoặc không thấy ô ký: heuristic cũ — trang chứa chức danh cuối cùng; trang sau còn chữ mà
    không mở đầu bằng Phụ lục/Biểu/Danh mục/Mẫu → nghi gãy.
    """
    gp = [_gon(pg) for pg in pages]
    dong = None
    if docx_path is not None:
        try:
            dong = _o_ky_docx(Path(docx_path))
        except Exception:
            dong = None
    if dong and len(dong) >= 2:
        ten, dau = dong[-1], dong[0]
        trang_ten = max((i for i, pg in enumerate(gp) if ten in pg), default=None)
        if trang_ten is not None:
            trang_dau = max((i for i, pg in enumerate(gp[:trang_ten + 1]) if dau in pg), default=None)
            if trang_dau is not None and trang_dau != trang_ten:
                # Cùng một ô bảng ký bị LibreOffice ngắt ngang: Word có thể ngắt khác (mẫu thật thong-bao-ket-luan-ubnd
                # của VP UBND tỉnh đã ban hành vẫn vậy, 18/10/2026) → qa_all.py hạ WARN kèm cách chữa H13(e) cantSplit.
                return [f"khối ký gãy trang trong cùng một ô bảng ký: '{dau}' ở trang {trang_dau + 1}, tên người ký "
                        f"'{ten}' ở trang {trang_ten + 1} — thêm <w:cantSplit/> cho hàng bảng ký (H13) hoặc rút thân"]
            return []
    if docx_path is not None and not dong:
        # có file .docx mà không thấy khối ký (phụ biểu, biên bản ký nhiều bên) — soát tay, không đoán theo trang
        return ["không tìm thấy khối ký trong file .docx (kiểm tra thủ công)"]
    issues = []
    title_page = None
    for pno, page in enumerate(pages, 1):
        for t in SIGN_TITLES:
            if t in page:
                title_page = max(title_page or 0, pno)
    if title_page is None:
        return ["không tìm thấy chức danh ký trong PDF (kiểm tra thủ công)"]
    for pno in range(title_page, len(pages)):
        body = pages[pno].strip()
        dau_trang = re.sub(r"^\d+\s+", "", " ".join(body.split()[:6]))
        if body and not re.match(r"(?i)(PHỤ\s+LỤC|Phụ\s+lục|BIỂU|Biểu|DANH\s+MỤC|Danh\s+mục|MẪU|Mẫu)", dau_trang):
            first_words = " ".join(body.split()[:12])
            issues.append(
                f"khối ký nghi gãy trang: chức danh ở trang {title_page}, "
                f"trang {pno + 1} vẫn còn nội dung: '{first_words}...'"
            )
    return issues


def check_line_shapes(docx_path: Path, min_lines: int) -> list:
    d = Document(str(docx_path))
    xml = d.element.xml
    # đếm cả VML (w:pict) lẫn DrawingML (w:drawing) — mẫu thật dùng cả hai kiểu
    n = xml.count("<w:pict") + xml.count("<w:drawing")
    if n < min_lines:
        return [f"chỉ có {n} shape Line trong file (yêu cầu >= {min_lines}: "
                "dưới tên cơ quan VÀ dưới Tiêu ngữ). Nghi mất Line do gán run.text "
                "vào run neo shape (Quy tắc 11)."]
    return []


def check_keepnext_audit(docx_path: Path) -> list:
    """Nhóm K1: liệt kê paragraph có w:keepNext KHÔNG giống đề mục -> WARN.
    Đề mục hợp lệ: I./II./III./IV./V. + tiêu đề ngắn "N. Về ..."/"N. Một số ...".
    Khoản nội dung đánh số bị gán keepNext là nguyên nhân gây trống nửa cuối trang trong Word."""
    import zipfile, re as _re, unicodedata
    issues = []
    try:
        xml = zipfile.ZipFile(docx_path).read('word/document.xml').decode('utf-8')
        for pm in _re.finditer(r'<w:p\b.*?</w:p>', xml, _re.S):
            seg = pm.group(0)
            if '<w:keepNext' not in seg:
                continue
            text = unicodedata.normalize('NFC', ''.join(_re.findall(r'<w:t[^>]*>([^<]*)</w:t>', seg))).strip()
            if not text:
                continue
            ok = bool(_re.match(r'^(?:[IVX]{1,4}\.|PHẦN|Phần|[0-9]{1,2}\.\s{0,2}(?:Về|Một số|Công tác|Kết quả|Đánh giá|Nhiệm vụ|Phương hướng))', text)) and len(text) <= 120
            if not ok:
                issues.append(f"keepNext trên đoạn KHÔNG giống đề mục (nguy cơ trống cuối trang - Nhóm K1): '{text[:70]}…'")
    except Exception as e:  # audit phụ, không làm hỏng QA chính
        issues.append(f"không audit được keepNext: {e}")
    return issues


def check_so_ngay_13pt(docx_path: Path) -> list:
    """Dòng 'Số: .../...' và dòng ngày trong bảng header: sz=26 tường minh,
    dòng ngày in nghiêng."""
    issues = []
    d = Document(str(docx_path))
    if not d.tables:
        return []
    t0 = d.tables[0]
    for r in range(len(t0.rows)):
        for c in range(len(t0.columns)):
            try:
                cell = t0.cell(r, c)
            except Exception:
                continue
            for p in cell.paragraphs:
                text = p.text
                is_so = bool(re.search(r"Số\s*:", text))
                is_ngay = bool(re.search(r"ngày\s.*tháng\s.*năm", text))
                if not (is_so or is_ngay):
                    continue
                for run in p.runs:
                    if not run.text.strip():
                        continue
                    rpr = run._element.find(qn("w:rPr"))
                    sz = rpr.find(qn("w:sz")) if rpr is not None else None
                    val = sz.get(qn("w:val")) if sz is not None else None
                    if val != "26":
                        issues.append(
                            f"dòng {'Số' if is_so else 'ngày'} run '{run.text[:20]}' "
                            f"sz={val} (yêu cầu 26 = 13pt tường minh)")
                        break
                if is_ngay:
                    has_italic = any(
                        run._element.find(qn("w:rPr")) is not None and
                        run._element.find(qn("w:rPr")).find(qn("w:i")) is not None
                        for run in p.runs if run.text.strip())
                    if not has_italic:
                        issues.append("dòng ngày thiếu in nghiêng (w:i)")
    return issues


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(2)
    docx_path = Path(args[0])
    # Guard đầu vào: tool nhận .docx (tự render PDF bên trong), KHÔNG nhận PDF.
    # Bài học 05/7/2026: truyền nhầm .pdf → PackageNotFoundError khó hiểu.
    if not docx_path.exists():
        print(f"LỖI: không tìm thấy file: {docx_path}")
        sys.exit(2)
    if docx_path.suffix.lower() != ".docx":
        print(f"LỖI: tool này nhận file .docx (tự render PDF để QA), "
              f"không nhận '{docx_path.suffix}'.")
        print(f"Cách dùng đúng:  python3 qa_pdf_check.py duong_dan/van_ban.docx")
        if docx_path.suffix.lower() == ".pdf":
            print("Gợi ý: nếu chỉ có PDF, dùng scripts/extract_metadata.py "
                  "hoặc pdftotext -layout để kiểm nội dung.")
        sys.exit(2)
    min_lines = 2
    do_render = "--no-render" not in args
    if "--min-lines" in args:
        min_lines = int(args[args.index("--min-lines") + 1])
    # --pdf <path>: dùng PDF đã render sẵn, KHÔNG render lại (tiết kiệm 1 lần soffice).
    pre_pdf = None
    if "--pdf" in args:
        pre_pdf = Path(args[args.index("--pdf") + 1])
        if not pre_pdf.exists():
            print(f"LỖI: --pdf không tìm thấy: {pre_pdf}")
            sys.exit(2)

    fails, warns = [], []

    # Kiểm tra XML (không cần render)
    for msg in check_line_shapes(docx_path, min_lines):
        fails.append(("LINES", msg))

    for msg in check_keepnext_audit(docx_path):
        warns.append(("KEEPNEXT", msg))
    for msg in check_so_ngay_13pt(docx_path):
        fails.append(("SZ13", msg))

    if do_render:
        with tempfile.TemporaryDirectory() as td:
            try:
                pdf = pre_pdf if pre_pdf else render_pdf(docx_path, Path(td))
                pages = pdf_pages_text(pdf)
            except Exception as e:
                warns.append(("RENDER", f"không render được PDF: {e}"))
                pages = None
            if pages:
                for msg in check_widow(pages):
                    fails.append(("WIDOW", msg))
                for msg in check_sig_split(pages, docx_path):
                    fails.append(("SIGSPLIT", msg))

    for tag, msg in warns:
        print(f"[WARN {tag}] {msg}")
    for tag, msg in fails:
        print(f"[FAIL {tag}] {msg}")
    if not fails:
        print("QA PASS: không phát hiện widow word, khối ký gãy trang, "
              "thiếu Line header, hay sai cỡ dòng Số/Ngày.")
        sys.exit(0)
    print(f"\nTổng: {len(fails)} lỗi — sửa xong chạy lại trước khi giao file.")
    sys.exit(1)


if __name__ == "__main__":
    main()
