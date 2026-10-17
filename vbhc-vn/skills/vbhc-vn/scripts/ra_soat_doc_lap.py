#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ra_soat_doc_lap.py — SỐ ĐO cho GATE cuối "Rà soát độc lập" (Bạn chốt 17/10/2026, bản giao việc mục 3.4).

Tác tử rà soát (chưa tham gia soạn) đọc file .docx BẰNG SCRIPT NÀY, không đọc từ trí nhớ phiên, rồi đối chiếu từng
dòng scripts/checklist_ra_soat.txt. Script chỉ ĐO và liệt kê, không sửa file, không kết luận thay người rà; dòng có
dấu "⚠" là chỗ lệch với checklist cần xem.

    python3 scripts/ra_soat_doc_lap.py file.docx            # bản đọc cho người/tác tử
    python3 scripts/ra_soat_doc_lap.py file.docx --json     # cho máy
    python3 scripts/ra_soat_doc_lap.py file.docx --pdf x.pdf   # dùng PDF đã render (kiem_tra_truoc_khi_giao.py truyền vào)

Các số đo: [1] phông từng run (kể cả kiểu, theme); [2] cỡ chữ khối Nơi nhận; [3] khoảng cách sau "Kính gửi" và trước
khối ký, số dòng trống ô ký so với số dòng Nơi nhận; [4] trang gần trống (cần PDF); [5] "UBND" trong thân; [6] định
dạng ngày tháng; [7] bảng: màu nền ô, căn dọc ô; [8] dấu vết chỉnh sửa; [9] dòng "Số:"; [10] Ban Quản lý Khu kinh tế
↔ Ban Quản lý các KCN; [11] tên cụm gộp số; [12] Nơi nhận văn bản thủ tục hành chính; [13] giao việc cho Phòng QLCN;
[14] chữ tím, chỗ trống; [15] "Kính thưa các đồng chí" giữa bài; [16] số hiệu văn bản quy phạm chưa có nguồn (R23).

Thoát mã 0 khi đo xong (kể cả có ⚠); 2 khi không đọc được file.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from qa_rules import (Ctx, _co_chu, _co_dau_doan, _la_tthc, _o_noi_nhan, _tt, cut, nfc,  # noqa: E402
                      rule_R23, text_runs)

PHONG_CHUAN = "Times New Roman"
TIM = {"7030A0", "7030a0"}
THANG = re.compile(r"ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(\d{4})", re.I)
NGAY_GACH = re.compile(r"(?<![\d/])(\d{1,2})/(\d{1,2})/(\d{4})(?![\d/])")
GOP_SO = re.compile(r"(?:Cụm công nghiệp|CCN|Khu công nghiệp|KCN)\s+[A-ZĐÂÊÔƠƯÁÀẢÃẠ][^,;.()]{0,40}?\s\d{1,2}\s*(?:,\s*\d{1,2}\s*)*(?:,|và)\s*\d{1,2}\b")
DAU_VET = re.compile(r"bản\s+cũ|bản\s+sửa|sửa\s+lần|dự\s+thảo\s+lần|\(cũ\)|[\w\-]+\.(?:docx|doc|pdf|xlsx)\b|\bv\d+\b|TODO|XXX", re.I)


def _ve(ds: list, muc: str, noi_dung: str, canh_bao: bool = False):
    ds.append({"muc": muc, "canh_bao": canh_bao, "noi_dung": noi_dung})


def _phong(path: Path) -> dict[str, int]:
    dem: dict[str, int] = {}
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if re.match(r"word/(document|styles|header\d*|footer\d*|footnotes|endnotes|numbering)\.xml$", n):
                x = z.read(n).decode("utf-8", "replace")
                for attrs in re.findall(r"<w:rFonts\b([^>]*)/>", x):
                    for v in set(re.findall(r'w:(?:ascii|hAnsi)="([^"]+)"', attrs)):
                        dem[v] = dem.get(v, 0) + 1
            elif n == "word/theme/theme1.xml":
                for v in re.findall(r'<a:latin typeface="([^"]*)"', z.read(n).decode("utf-8", "replace")):
                    if v:
                        dem[f"{v} (theme)"] = dem.get(f"{v} (theme)", 0) + 1
    return dem


def _trang_pdf(pdf: Path) -> list[str]:
    r = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, timeout=60)
    trang = r.stdout.split("\f")
    if trang and not trang[-1].strip():
        trang.pop()
    return trang


def _render(docx: Path) -> Path | None:
    qa = Path("/home/claude/work/qa") / (docx.stem + ".pdf")
    if qa.exists() and qa.stat().st_mtime >= docx.stat().st_mtime:
        return qa  # qa_all.py vừa render, dùng lại
    out = Path(tempfile.mkdtemp(prefix="rasoat-"))
    try:
        subprocess.run(["soffice", "--headless", "-env:UserInstallation=file:///home/claude/.lo_profile",
                        "--convert-to", "pdf", "--outdir", str(out), str(docx)],
                       check=True, capture_output=True, timeout=180)
    except Exception:
        return None
    p = out / (docx.stem + ".pdf")
    return p if p.exists() else None


def do(path: Path, pdf: Path | None = None, render: bool = True) -> list[dict]:
    doc = Document(str(path))
    ctx = Ctx(path)
    ctx.build(doc)
    ds: list[dict] = []
    than = [(loc, t) for loc, t in ctx.texts if loc.startswith("đoạn")]

    # [1] Phông
    ph = _phong(path)
    la = {k: v for k, v in ph.items() if not k.startswith(PHONG_CHUAN) and k not in ("Symbol", "Wingdings")}
    _ve(ds, "1. Phông", "; ".join(f"{k}: {v}" for k, v in sorted(ph.items())) or "không khai báo phông riêng",
        bool(la))

    # [2] Nơi nhận
    _, o = _o_noi_nhan(doc)
    so_dong_nn = 0
    if o is None:
        _ve(ds, "2. Nơi nhận", "không thấy ô 'Nơi nhận:'")
    else:
        for p in o.paragraphs:
            tx = nfc(p.text).strip()
            if not tx:
                continue
            so_dong_nn += 1
            co = sorted({c for c in (_co_chu(r, p, doc) for r in text_runs(p)) if c is not None})
            chuan = 12 if tx.startswith("Nơi nhận") else 11
            _ve(ds, "2. Nơi nhận", f"cỡ {', '.join(f'{c:g}' for c in co) or '?'} — {cut(tx, 60)}",
                any(abs(c - chuan) > 0.01 for c in co))

    # [3] Khoảng cách Kính gửi, khối ký
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if nfc(p.text).strip().startswith("Kính gửi"):
            j = i + 1
            while j < len(paras) and nfc(paras[j].text).strip().startswith(("-", "+")) and paras[j].text.strip():
                j += 1
            trong = 0
            while j < len(paras) and not paras[j].text.strip():
                trong += 1
                j += 1
            _ve(ds, "3. Sau Kính gửi", f"{trong} dòng trống rồi mới vào thân (chuẩn: 1)", trong != 1)
            break
    t, _ = _o_noi_nhan(doc)
    if t is not None and t._tbl.getparent() is doc.element.body:
        kids = list(doc.element.body)
        k = kids.index(t._tbl) - 1
        trong = []
        while k >= 0 and kids[k].tag == qn("w:p") and not _tt(kids[k]).strip():
            trong.append(kids[k])
            k -= 1
        _ve(ds, "3. Trước khối ký", f"{len(trong)} dòng trống giữa thân và khối Nơi nhận - ký (chuẩn: 1, cỡ 14)",
            len(trong) != 1)
        o_ky = None
        for row in t.rows:
            for c in row.cells:
                if c is not o and re.search(r"GIÁM ĐỐC|CHỦ TỊCH|TRƯỞNG|KT\.|TM\.|Q\.", nfc(c.text)):
                    o_ky = c
        if o_ky is not None:
            pp = o_ky.paragraphs
            chu = [i for i, p in enumerate(pp) if p.text.strip()]
            giua = 0
            if len(chu) >= 2:
                giua = sum(1 for p in pp[chu[-2] + 1:chu[-1]] if not p.text.strip())
            rong = o_ky._tc.find(qn("w:tcPr"))
            w = rong.find(qn("w:tcW")).get(qn("w:w")) if rong is not None and rong.find(qn("w:tcW")) is not None else "?"
            _ve(ds, "3. Ô ký", f"{giua} dòng trống giữa chức danh và tên người ký; ô Nơi nhận {so_dong_nn} dòng; "
                f"ô ký rộng {w} dxa (chuẩn: ≥ 3 dòng trống, cao tương ứng khối Nơi nhận)",
                giua < 3 or (so_dong_nn and giua + 3 < so_dong_nn - 2))

    # [4] Trang gần trống
    if pdf is None and render:
        pdf = _render(path)
    if pdf is not None:
        trang = _trang_pdf(pdf)
        for i, tx in enumerate(trang, 1):
            n = len(re.sub(r"\s", "", tx))
            if n < 120:
                _ve(ds, "4. Trang", f"trang {i}/{len(trang)} chỉ có {n} ký tự — trang trống hoặc khoảng trắng lớn", True)
        if not any(x["muc"] == "4. Trang" for x in ds):
            _ve(ds, "4. Trang", f"{len(trang)} trang, không trang nào gần trống")
    else:
        _ve(ds, "4. Trang", "không render được PDF — xem ảnh qa_all.py để soát trang trống")

    # [5] UBND trong thân (ngoài bảng header, Nơi nhận)
    ubnd = [(loc, t) for loc, t in than if re.search(r"\bUBND\b", t)]
    for loc, t in ubnd[:10]:
        _ve(ds, "5. UBND viết tắt", f"{loc}: {cut(t, 70)}", True)
    if not ubnd:
        _ve(ds, "5. UBND viết tắt", "0 lần trong thân văn bản")

    # [6] Ngày tháng
    sai_ngay = []
    for loc, t in ctx.texts:
        for m in THANG.finditer(t):
            d, mo = m.group(1), m.group(2)
            if (int(mo) in (1, 2) and len(mo) == 1) or (int(mo) >= 3 and mo.startswith("0")) \
                    or (int(d) < 10 and len(d) == 1):
                sai_ngay.append((loc, m.group(0)))
        for m in NGAY_GACH.finditer(t):
            mo = m.group(2)
            if (int(mo) in (1, 2) and len(mo) == 1) or (int(mo) >= 3 and mo.startswith("0")):
                sai_ngay.append((loc, m.group(0)))
    for loc, x in sai_ngay[:15]:
        _ve(ds, "6. Ngày tháng", f"{loc}: '{x}' (ngày < 10 và tháng 01, 02 có số 0; tháng 3–12 không có số 0)", True)
    if not sai_ngay:
        _ve(ds, "6. Ngày tháng", "mọi ngày tháng đúng quy ước số 0")

    # [7] Bảng nội dung: nền, căn dọc (bỏ bảng header và bảng ký)
    for ti, tb in enumerate(doc.tables):
        if tb is t or ti == 0:
            continue
        nen, doc_ = set(), set()
        for tc in tb._tbl.iter(qn("w:tc")):
            pr = tc.find(qn("w:tcPr"))
            shd = pr.find(qn("w:shd")) if pr is not None else None
            fill = shd.get(qn("w:fill")) if shd is not None else None
            if fill and fill.upper() not in ("FFFFFF", "AUTO"):
                nen.add(fill)
            va = pr.find(qn("w:vAlign")) if pr is not None else None
            doc_.add(va.get(qn("w:val")) if va is not None else "top (mặc định)")
        _ve(ds, "7. Bảng", f"bảng {ti}: màu nền {', '.join(sorted(nen)) or 'trắng'}; căn dọc {', '.join(sorted(doc_))}",
            bool(nen) or doc_ != {"center"})

    # [8] Dấu vết chỉnh sửa
    xml = ctx.xml
    n_ins, n_del = xml.count("<w:ins "), xml.count("<w:del ")
    n_cmt = ctx.comments_xml.count("<w:comment ")
    _ve(ds, "8. Dấu vết", f"theo dõi thay đổi: {n_ins} chèn, {n_del} xóa; {n_cmt} chú thích", bool(n_ins or n_del or n_cmt))
    for loc, t in ctx.texts:
        for m in DAU_VET.finditer(t):
            _ve(ds, "8. Dấu vết", f"{loc}: '{m.group(0)}' — {cut(t, 60)}", True)
    dc = [(loc, t) for loc, t in than if re.search(r"điều\s+chỉnh", t, re.I)]
    if dc:
        _ve(ds, "8. Dấu vết", f"'điều chỉnh' xuất hiện {len(dc)} lần — soát bằng mắt (có thể hợp lệ): "
            + "; ".join(loc for loc, _ in dc[:5]))

    # [9] Dòng Số
    for loc, t in ctx.texts[:30]:
        if re.match(r"^\s*Số\s*:", t):
            so = re.sub(r"^\s*Số\s*:\s*", "", t).strip()
            _ve(ds, "9. Số văn bản", f"'{cut(t, 50)}' — {'để trống số' if re.match(r'^[/\s]', so) else 'ĐÃ có số'}"
                " (văn bản soạn cho doanh nghiệp, cơ quan khác: để trống số)")
            break

    # [10] Ban Quản lý Khu kinh tế
    if re.search(r"Ban\s+Quản\s+lý\s+Khu\s+kinh\s+tế", ctx.full_text) and \
            not re.search(r"Ban\s+Quản\s+lý\s+các\s+Khu\s+công\s+nghiệp", ctx.full_text):
        _ve(ds, "10. Ban Quản lý", "có 'Ban Quản lý Khu kinh tế' nhưng không có 'Ban Quản lý các Khu công nghiệp tỉnh'", True)

    # [11] Tên cụm gộp số
    for loc, t in ctx.texts:
        for m in GOP_SO.finditer(t):
            _ve(ds, "11. Tên cụm", f"{loc}: '{m.group(0)}' — viết đủ tên từng cụm", True)

    # [12] Văn bản thủ tục hành chính gửi Ban Giám đốc Sở
    if o is not None:
        tthc = _la_tthc(ctx)
        bgd = "Ban Giám đốc" in nfc(o.text)
        _ve(ds, "12. Nơi nhận TTHC", f"văn bản thủ tục hành chính: {'có' if tthc else 'không'}; "
            f"Nơi nhận có Ban Giám đốc Sở: {'có' if bgd else 'không'}", tthc and bgd)

    # [13] Giao việc cho Phòng Quản lý Công nghiệp trong văn bản do Phòng tham mưu
    luu_cn = re.search(r"Lưu\s*:\s*VT\s*,\s*CN\b", ctx.full_text)
    for loc, t in than:
        if luu_cn and re.search(r"(giao|đề nghị|yêu cầu)[^.;]{0,40}Phòng\s+Quản\s+lý\s+Công\s+nghiệp", t, re.I):
            _ve(ds, "13. Giao việc", f"{loc}: {cut(t, 70)} — văn bản Phòng QLCN tham mưu không giao việc cho chính Phòng", True)

    # [14] Chữ tím, chỗ trống
    tim = sum(1 for p in doc.paragraphs for r in p.runs
              if r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb) in TIM)
    trong = [(loc, t) for loc, t in than if re.search(r"…{2,}|\.{4,}|\[\s*\]|\bxx+\b", t)]
    _ve(ds, "14. Chỗ chưa hoàn thiện", f"{tim} run chữ tím; {len(trong)} đoạn còn chỗ trống "
        + ("(" + ", ".join(loc for loc, _ in trong[:5]) + ")" if trong else ""), bool(trong))

    # [15] Kính thưa giữa bài
    kt = [(i, p) for i, p in enumerate(paras) if nfc(p.text).strip().startswith("Kính thưa các đồng chí")]
    for i, p in kt[1:]:
        ind = p.paragraph_format.first_line_indent
        _ve(ds, "15. Kính thưa", f"đoạn {i}: lùi đầu dòng {ind.cm if ind is not None else 0:.2f} cm",
            ind is None or ind.cm < 0.5)

    # [16] Số hiệu văn bản quy phạm chưa có nguồn
    r23 = rule_R23(doc, ctx)
    for f in r23:
        so = re.search(r"'([^']+)'", f.hint).group(1)
        _ve(ds, "16. Căn cứ chưa có nguồn", f"{so} ({f.loc})", True)
    if not r23:
        _ve(ds, "16. Căn cứ chưa có nguồn", "mọi số hiệu văn bản quy phạm đều khớp kho hoặc căn cứ của phiên")
    return ds


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("--pdf", help="PDF đã render của chính file này")
    ap.add_argument("--khong-render", action="store_true", help="bỏ đo trang trống (không gọi LibreOffice)")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    p = Path(a.docx)
    try:
        ds = do(p, Path(a.pdf) if a.pdf else None, render=not a.khong_render)
    except Exception as e:
        print(f"LỖI đọc {p}: {type(e).__name__}: {e}", file=sys.stderr)
        return 2
    if a.json:
        print(json.dumps(ds, ensure_ascii=False, indent=2))
        return 0
    print("=" * 66)
    print(f"SỐ ĐO RÀ SOÁT ĐỘC LẬP — {p.name}")
    print("=" * 66)
    for x in ds:
        print(f"{'⚠' if x['canh_bao'] else ' '} [{x['muc']}] {x['noi_dung']}")
    n = sum(1 for x in ds if x["canh_bao"])
    print(f"\n{n} chỗ lệch cần đối chiếu checklist_ra_soat.txt.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
